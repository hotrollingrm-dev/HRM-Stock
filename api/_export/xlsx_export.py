"""
COPY of backend/export.py (Excel part) for the Vercel export function — keep in sync.
Export the monthly A3 stock sheets.

  • Excel — the original "Stock RM" workbook layout: one sheet per month (Jan, Feb, …),
    blocks of 6 materials × 31 days, same fonts / colours / column widths / formulas
    (ทั้งหมด = คงเหลือวันก่อน, คงเหลือ = ทั้งหมด + รับสินค้า − ใช้ไป, SUM row, OPS linked to the first column).
    Styles come from backend/templates/a3_template.xlsx (built from the original file).
  • PDF   — the same A3 HTML the app prints, rendered by Edge/Chrome headless (420 × 297 mm).
"""
import copy
import glob
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "a3_template.xlsx")

MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
SHEET = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
UNIT_LONG = {"L": "Liter", "Pail": "Pails", "Drum": "Drums", "Roll": "Rolls", "Box": "Box"}
BLOCK = 37            # rows per block (header … sum row + 1 spacer)
FIRST = 2             # header row of block 1
COLS = (2, 8, 14, 20, 26, 32)   # B H N T Z AF


def _n(v):
    if v is None:
        return ""
    v = float(v)
    return f"{int(v):,}" if v == int(v) else f"{v:,.2f}".rstrip("0").rstrip(".")


def _unit(u, q):
    if u in UNIT_LONG:
        return u if (q == 1 and u != "L") else UNIT_LONG[u]
    return u or ""


def _num(v):
    if v is None:
        return None
    return int(v) if float(v) == int(v) else float(v)


def _copy_rows(src, dst, r0, r1, offset):
    """Copy styles + heights + merges of src rows r0..r1 to dst rows shifted by offset (same workbook)."""
    for r in range(r0, r1 + 1):
        h = src.row_dimensions[r].height
        if h:
            dst.row_dimensions[r + offset].height = h
        for c in range(1, 37):
            s = src.cell(r, c)
            if s.has_style:
                dst.cell(r + offset, c)._style = copy.copy(s._style)
    for rng in src.merged_cells.ranges:
        if r0 <= rng.min_row and rng.max_row <= r1:
            dst.merge_cells(start_row=rng.min_row + offset, start_column=rng.min_col,
                            end_row=rng.max_row + offset, end_column=rng.max_col)


def _col_widths(tpl):
    out = {}
    for dim in tpl.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for c in range(dim.min, dim.max + 1):
                out[c] = dim.width
    return {c: w for c, w in out.items() if c <= 36}


def _print_scale(tpl):
    """Largest whole % at which row 1 + one 37-row block fits A3 landscape (1190 × 842 pt, small margins)."""
    cw = _col_widths(tpl)
    width_pt = sum((cw.get(c, 9.14) * 7 + 5) * 0.75 for c in range(1, 37))
    height_pt = sum(tpl.row_dimensions[r].height or 15 for r in range(1, FIRST + BLOCK - 1))
    s = min((1190.55 - 0.5 * 72) / width_pt, (841.89 - 0.6 * 72) / height_pt)
    return max(10, min(100, int(s * 100) - 1))


def _sheet_name(month, multi_year):
    y, m = int(month[:4]), int(month[5:7])
    return f"{SHEET[m - 1]} {y}" if multi_year else SHEET[m - 1]


def write_xlsx(path, months_data, blank=False, title=""):
    """months_data: list of api.month_get() results (in order)."""
    wb = openpyxl.load_workbook(TEMPLATE)
    tpl = wb["TEMPLATE"]
    multi_year = len({m["month"][:4] for m in months_data}) > 1
    for mg in months_data:
        month, nd = mg["month"], mg["days"]
        y, mo = int(month[:4]), int(month[5:7])
        mats = mg["materials"]
        nblocks = max(1, (len(mats) + 5) // 6)
        ws = wb.create_sheet(_sheet_name(month, multi_year))
        for c, w in _col_widths(tpl).items():            # template stores ranges (C:D, E:F …) — expand them
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.page_setup.paperSize = 8                   # A3
        ws.page_setup.orientation = "landscape"
        # one block per A3 sheet: fixed scale + manual page breaks (Excel ignores breaks when "Fit to" is on)
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0, footer=0)
        ws.page_setup.scale = _print_scale(tpl)
        ws.print_options.horizontalCentered = True
        ws.sheet_view.zoomScale = 70

        _copy_rows(tpl, ws, 1, 1, 0)
        ws["B1"] = "UPDATE :"
        if not blank:
            ws["D1"] = "=TODAY()"
        if title:                                     # optional heading, centred over the sheet (e.g. "Stock HRM")
            ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=25)
            c = ws.cell(1, 12, title)
            c.font = Font(bold=True, size=18)
            c.alignment = Alignment(horizontal="center", vertical="center")
        label = f"{MON[mo - 1]} {y}"
        ops_src = None  # OPS is typed once (first material) — the other blocks link to it, like the original
        for b in range(nblocks):
            h = FIRST + b * BLOCK
            # header lines, sum row, spacer come straight from the template block
            _copy_rows(tpl, ws, FIRST, FIRST + 3, h - FIRST)
            _copy_rows(tpl, ws, FIRST + 35, FIRST + 36, h - FIRST)
            # day rows by role, like the original sheets: day 1 (opening in red) · normal day ·
            # last day of month (balance highlighted yellow) · rows after month-end (grey)
            for d in range(1, 32):
                src = 6 if d == 1 else 35 if d == nd else 7 if d < nd else 36
                _copy_rows(tpl, ws, src, src, h + 3 + d - src)
            ws.cell(h, 1, label)
            for d in range(1, nd + 1):                # day numbers (column A) — independent of the slots
                ws.cell(h + 3 + d, 1, d)
            for i, c in enumerate(COLS):
                ws.cell(h + 3, c, "รับสินค้า")
                ws.cell(h + 3, c + 1, "ทั้งหมด")
                ws.cell(h + 3, c + 2, "ใช้ไป")
                ws.cell(h + 3, c + 3, "คงเหลือ")
                ws.cell(h + 3, c + 4, "OPS")
                idx = b * 6 + i
                m = mats[idx] if idx < len(mats) else None
                if m:
                    ws.cell(h, c, m["name"])
                    col = (m.get("a3_color") or "FFCC99").upper()
                    ws.cell(h, c).fill = PatternFill("solid", fgColor="FF" + col)
                    ws.cell(h + 1, c, f"จำนวนคงเหลือที่ต้องสั่ง : {_n(m['reorder_point'])} {_unit(m['unit'], m['reorder_point'])}")
                    moq = " (MOQ)" if m.get("moq") and m["moq"] == m["order_qty"] else ""
                    ws.cell(h + 2, c, f"จำนวนที่สั่งซื้อ : {_n(m['order_qty'])} {_unit(m['unit'], m['order_qty'])}{moq}")
                else:                                 # no material in this slot → leave it completely empty
                    for rng in [g for g in ws.merged_cells.ranges
                                if h <= g.min_row and g.max_row <= h + 35 and c <= g.min_col and g.max_col <= c + 4]:
                        ws.unmerge_cells(str(rng))
                    for r in range(h, h + 36):
                        for k in range(5):
                            ws.cell(r, c + k).style = "Normal"
                            ws.cell(r, c + k).value = None
                    continue
                for d in range(1, 32):
                    r = h + 3 + d
                    if d > nd:
                        continue
                    if not m or blank:
                        continue
                    L = get_column_letter
                    rc, tc, uc, bc, oc = (L(c + k) for k in range(5))
                    x = m["days"][d - 1]
                    if x["rcv"] is not None:
                        ws.cell(r, c, _num(x["rcv"]))
                    ws.cell(r, c + 1, _num(m["opening"]) or 0 if d == 1 else f"={bc}{r - 1}")
                    if x["used"] is not None:
                        ws.cell(r, c + 2, _num(x["used"]))
                    ws.cell(r, c + 3, f"={tc}{r}+{rc}{r}-{uc}{r}")
                    if ops_src is None:
                        ops_src = oc
                    if ops_src == oc and b == 0:
                        if mg["ops"].get(d) or mg["ops"].get(str(d)):
                            ws.cell(r, c + 4, mg["ops"].get(d) or mg["ops"].get(str(d)))
                    else:
                        ws.cell(r, c + 4, f'=IF({ops_src}{FIRST + 3 + d}="","",{ops_src}{FIRST + 3 + d})')
                if m and not blank:
                    s = h + 35
                    first, lastr = h + 4, h + 3 + nd
                    ws.cell(s, c, f"=SUM({L(c)}{first}:{L(c)}{lastr})")
                    ws.cell(s, c + 2, f"=SUM({L(c + 2)}{first}:{L(c + 2)}{lastr})")
            rows = list(range(h, h + 36)) + ([1] if b == 0 else [])
            for r in rows:                    # never show "####" when a number is wider than its column
                for c in range(1 if r > h + 3 or r == 1 else 2, 37):
                    cell = ws.cell(r, c)
                    al = cell.alignment
                    if cell.has_style and not al.shrink_to_fit:
                        cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or "center",
                                                   text_rotation=al.text_rotation, indent=al.indent, shrink_to_fit=True)
            if b < nblocks - 1:
                ws.row_breaks.append(Break(id=h + BLOCK - 1))
        ws.print_area = f"A1:AJ{FIRST + nblocks * BLOCK - 2}"
    wb.remove(tpl)
    wb.active = 0
    wb.save(path)
    return path


# ---------------------------------------------------------------- PDF
